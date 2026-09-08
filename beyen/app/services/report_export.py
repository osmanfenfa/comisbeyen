"""
Report export utilities.
export_to_pdf  — returns bytes (PDF)
export_to_excel — returns bytes (XLSX)
"""
import io
from typing import List, Dict, Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def export_to_pdf(data: List[Dict[str, Any]], title: str, columns: List[str]) -> bytes:
    """
    Render a tabular PDF report.
    data    — list of row dicts
    title   — report heading
    columns — ordered list of dict keys to include as columns
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1 * cm, leftMargin=1 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"<b>{title}</b>", styles["Title"]))
    elements.append(Spacer(1, 0.5 * cm))

    # Table header + rows
    header = [col.replace("_", " ").title() for col in columns]
    rows = [[str(row.get(col, "")) for col in columns] for row in data]
    table_data = [header] + rows

    col_count = len(columns)
    available_width = landscape(A4)[0] - 2 * cm
    col_width = available_width / col_count

    table = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F8E9")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.read()


def export_to_excel(data: List[Dict[str, Any]], title: str, columns: List[str]) -> bytes:
    """Render a tabular XLSX report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel sheet name limit

    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    alt_fill = PatternFill(start_color="F1F8E9", end_color="F1F8E9", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    # Header row
    headers = [col.replace("_", " ").title() for col in columns]
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    # Data rows
    for r_idx, row in enumerate(data, 2):
        fill = alt_fill if r_idx % 2 == 0 else None
        for c_idx, col in enumerate(columns, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row.get(col, ""))
            if fill:
                cell.fill = fill

    # Auto-size columns
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
